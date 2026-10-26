import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_border(cell, **kwargs):
    """
    Set cell borders
    kwargs: top, bottom, left, right
    values: dict(sz=12, val='single', color='000000', space='0')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(r'''
        <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:val="{top_val}" w:sz="{top_sz}" w:space="0" w:color="{top_color}"/>
            <w:left w:val="{left_val}" w:sz="{left_sz}" w:space="0" w:color="{left_color}"/>
            <w:bottom w:val="{bottom_val}" w:sz="{bottom_sz}" w:space="0" w:color="{bottom_color}"/>
            <w:right w:val="{right_val}" w:sz="{right_sz}" w:space="0" w:color="{right_color}"/>
        </w:tcBorders>
    '''.format(
        top_val=kwargs.get('top', {}).get('val', 'single'),
        top_sz=kwargs.get('top', {}).get('sz', '4'),
        top_color=kwargs.get('top', {}).get('color', '000000'),
        left_val=kwargs.get('left', {}).get('val', 'single'),
        left_sz=kwargs.get('left', {}).get('sz', '4'),
        left_color=kwargs.get('left', {}).get('color', '000000'),
        bottom_val=kwargs.get('bottom', {}).get('val', 'single'),
        bottom_sz=kwargs.get('bottom', {}).get('sz', '4'),
        bottom_color=kwargs.get('bottom', {}).get('color', '000000'),
        right_val=kwargs.get('right', {}).get('val', 'single'),
        right_sz=kwargs.get('right', {}).get('sz', '4'),
        right_color=kwargs.get('right', {}).get('color', '000000'),
    ))
    tcPr.append(tcBorders)

def add_placeholder_box(doc, figure_num, title, description):
    """Creates a clean black-and-white empty box for screenshots with a label."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.2)
    
    # Border
    border_args = {'val': 'single', 'sz': '8', 'color': '000000'}
    set_cell_border(cell, top=border_args, bottom=border_args, left=border_args, right=border_args)
    
    # Content inside placeholder box
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(12)
    
    run1 = p.add_run(f"[INSERT SCREENSHOT HERE: {title.upper()}]\n")
    run1.bold = True
    run1.font.size = Pt(11)
    run1.font.color.rgb = RGBColor(0, 0, 0)
    
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(36)
    
    run2 = p2.add_run(f"Recommended Content: {description}")
    run2.font.size = Pt(9.5)
    run2.font.italic = True
    run2.font.color.rgb = RGBColor(60, 60, 60)
    
    # Caption below box
    cap_p = doc.add_paragraph()
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_p.paragraph_format.space_before = Pt(6)
    cap_p.paragraph_format.space_after = Pt(14)
    cap_run = cap_p.add_run(f"Figure {figure_num}: {title}")
    cap_run.bold = True
    cap_run.font.size = Pt(10)
    cap_run.font.color.rgb = RGBColor(0, 0, 0)

print("Helper functions ready.")
